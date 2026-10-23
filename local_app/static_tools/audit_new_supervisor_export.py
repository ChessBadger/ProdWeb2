"""Read-only, all-account supervisor-production audit of the original workbook."""
import collections
import datetime
import hashlib
import html
import json
import math
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
SOURCE = Path('C:/Users/Laptop 122/Desktop/Store Prep/EmployeeProductionExportLashaunfull.xlsx')
OUT = ROOT / 'local_app/verification/new-supervisor-export-2026-10-15'
OUT.mkdir(exist_ok=True)

def ident(value):
    s = str(value or '').strip()
    if s.endswith('.0'): s = s[:-2]
    return s.zfill(4) if s and s != '0' else ''

def numeric(value):
    if value is None or not str(value).strip(): return None
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (TypeError, ValueError): return None

wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb.worksheets[0]
it = ws.iter_rows(values_only=True)
headers = next(it)
rows = [dict(zip(headers, values), excelRow=i) for i, values in enumerate(it, 2)]
names = {ident(r['Employee']): (str(r['FirstName'] or '')+' '+str(r['LastName'] or '')).strip() for r in rows}
visits = collections.defaultdict(list)
for r in rows:
    date = r['DateOfInv'].date().isoformat() if isinstance(r['DateOfInv'], datetime.datetime) else str(r['DateOfInv']).strip()[:10]
    visits[(str(r['AccountName'] or '').strip(), str(r['StoreName'] or '').strip(), date)].append(r)

detail = []
unassigned = []
for (account, store, date), rs in sorted(visits.items()):
    supervisors = sorted({ident(r['SupervisorNumber']) for r in rs} - {''})
    if not supervisors:
        unassigned.append(dict(account=account, store=store, date=date, rows=[r['excelRow'] for r in rs]))
    for sup in supervisors:
        own = [r for r in rs if ident(r['Employee']) == sup]
        qty = [numeric(r['Total_Ext_Qty']) for r in own]
        pph = [numeric(r['PiecesPerHr']) for r in own]
        status = ('No employee row' if not own else 'Positive pieces' if any(q is not None and q > 0 for q in qty)
                  else 'Negative pieces' if any(q is not None and q < 0 for q in qty)
                  else 'Zero pieces' if any(q == 0 for q in qty) else 'Blank/invalid pieces')
        detail.append(dict(account=account, store=store, date=date, supervisorId=sup,
            supervisorName=names.get(sup, 'Not in employee rows'), status=status,
            supervisorRowCount=len(own), productionQuantities=qty, piecesPerHour=pph,
            positiveRate=any(v is not None and v > 0 for v in pph),
            ownSourceRows=[r['excelRow'] for r in own], assignmentSourceRows=[r['excelRow'] for r in rs if ident(r['SupervisorNumber']) == sup],
            conflictingSupervisorIds=supervisors if len(supervisors) > 1 else [],
            ownRowAssignedToOtherSupervisor=any(ident(r['SupervisorNumber']) != sup for r in own)))

def summarize(records):
    c = collections.Counter(r['status'] for r in records)
    return dict(assignments=len(records), positivePieces=c['Positive pieces'], zeroPieces=c['Zero pieces'],
        negativePieces=c['Negative pieces'], blankInvalidPieces=c['Blank/invalid pieces'], missingEmployeeRow=c['No employee row'],
        presentEmployeeRow=sum(r['supervisorRowCount'] > 0 for r in records), positiveRate=sum(r['positiveRate'] for r in records),
        positiveCoverage=round(c['Positive pieces']/len(records)*100, 2) if records else None,
        conflictingAssignments=sum(bool(r['conflictingSupervisorIds']) for r in records), duplicateOwnRows=sum(r['supervisorRowCount'] > 1 for r in records))

accounts = [dict(account=a, **summarize([r for r in detail if r['account'] == a])) for a in sorted({r['account'] for r in detail})]
sup_accounts = [dict(account=a, supervisorId=s, supervisorName=names.get(s,'Not in employee rows'), **summarize([r for r in detail if r['account']==a and r['supervisorId']==s])) for a,s in sorted({(r['account'],r['supervisorId']) for r in detail})]
summary = dict(source=str(SOURCE), sourceSHA256=hashlib.sha256(SOURCE.read_bytes()).hexdigest(), sheet=ws.title,
    sourceRows=len(rows), inventoryCount=len(visits), distinctSupervisorIds=len({r['supervisorId'] for r in detail}),
    dateStart=min(k[2] for k in visits), dateEnd=max(k[2] for k in visits),
    method='All export rows, with no date, TypeOfInv, counting-mode, active-roster, or account filter. Inventory identity is exact trimmed account + store + date; each distinct assigned supervisor checked against Employee on that inventory. Production is Total_Ext_Qty; numeric zero is explicitly separate from missing.',
    totals=summarize(detail), unassignedInventories=unassigned, accounts=accounts, supervisorAccounts=sup_accounts)
assert sum(a['assignments'] for a in accounts) == len(detail)
assert sum(summary['totals'][k] for k in ['positivePieces','zeroPieces','negativePieces','blankInvalidPieces','missingEmployeeRow']) == len(detail)
assert len({(r['account'],r['store'],r['date'],r['supervisorId']) for r in detail}) == len(detail)
for r in detail:
    assert (r['status'] == 'No employee row') == (not r['ownSourceRows'])
(OUT/'summary.json').write_text(json.dumps(summary,indent=2),encoding='utf8')
(OUT/'inventory-detail.private.json').write_text(json.dumps(detail,indent=2),encoding='utf8')


import sys
sys.path.insert(0,str(ROOT))
from export_history import row_key,latest_payload,load_manifest
archive=Path.home()/'AppData/Local/EmployeePredictor/ExportHistory/d27543e5831b'
retained=latest_payload(archive,load_manifest(archive))
old={row_key(r):r for r in retained['EmployeeProductionHistory']}
new={};duplicates=[];conflicts=[];invalid=[];changed=collections.Counter()
for r in rows:
 clean={k:(v.isoformat() if isinstance(v,datetime.datetime) else v) for k,v in r.items() if k!='excelRow'}
 try:k=row_key(clean)
 except Exception as e:invalid.append([r['excelRow'],str(e)]);continue
 if k in new:
  duplicates.append(r['excelRow'])
  if new[k]!=clean:conflicts.append(r['excelRow'])
 new[k]=clean
 if k in old:
  for f,v in clean.items():
   ov=old[k].get(f)
   if f=='DateOfInv':v=str(v)[:10];ov=str(ov)[:10]
   if v!=ov:changed[f]+=1
active={ident(i) for i in retained['_meta']['latestEmployeeIds']}
newids={ident(r['Employee']) for r in rows}
checks={}
for field in ['Total_Ext_Qty','PiecesPerHr','Expr1','DollarPerHr','SkusPerHr']:
 vals=[numeric(r.get(field)) for r in rows]
 checks[field]={'blankInvalid':sum(v is None for v in vals),'zero':sum(v==0 for v in vals),'negative':sum(v is not None and v<0 for v in vals),'max':max((v for v in vals if v is not None),default=None)}
oldDetail=json.loads((ROOT/'local_app/verification/all-supervisors-2026-10-14/inventory-detail.private.json').read_text())
key=lambda r:(r['account'],r['store'],r['date'],r['supervisorId'])
newDetail={key(r):r for r in detail}
comparison=collections.Counter(newDetail.get(key(r),{}).get('status','Visit absent') for r in oldDetail if r['status']=='No employee row')
badHours=[d for d in detail if d['status']=='Positive pieces' and not any((numeric(r.get('Expr1')) or 0)>0 for r in visits[(d['account'],d['store'],d['date'])] if ident(r['Employee'])==d['supervisorId'])]
extra=dict(sheets=wb.sheetnames,headers=headers,duplicateRows=duplicates,conflictingRows=conflicts,invalidKeys=invalid,numericChecks=checks,previousMissingNow=dict(comparison),newKeys=len(new.keys()-old.keys()),retainedKeysAbsent=len(old.keys()-new.keys()),changedFields=dict(changed),currentRoster=len(active),newRoster=len(newids),rosterAdded=[dict(id=i,name=names.get(i)) for i in sorted(newids-active)],rosterMissing=sorted(active-newids),supervisorPositiveWithoutHours=badHours)
(OUT/'quality-comparison.json').write_text(json.dumps(extra,indent=2,default=str),encoding='utf8')
print(json.dumps({k:v for k,v in summary.items() if k not in ['accounts','supervisorAccounts','unassignedInventories']},indent=2))
print('ACCOUNTS',json.dumps(sorted(accounts,key=lambda x:-x['assignments'])[:18]))
print('QUALITY',json.dumps(extra,default=str))
print('MISSING',json.dumps([d for d in detail if d['status']!='Positive pieces']))
